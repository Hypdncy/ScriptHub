#!/usr/bin/env python3
# -*- coding:utf-8 -*-
# ------------------------------------------------------------
# File: base.py
# Created Date: 2020/6/24
# Created Time: 0:14
# Author: Hypdncy
# Author Mail: hypdncy@outlook.com
# Copyright (c) 2020 Hypdncy
# ------------------------------------------------------------
#                       .::::.
#                     .::::::::.
#                    :::::::::::
#                 ..:::::::::::'
#              '::::::::::::'
#                .::::::::::
#           '::::::::::::::..
#                ..::::::::::::.
#              ``::::::::::::::::
#               ::::``:::::::::'        .:::.
#              ::::'   ':::::'       .::::::::.
#            .::::'      ::::     .:::::::'::::.
#           .:::'       :::::  .:::::::::' ':::::.
#          .::'        :::::.:::::::::'      ':::::.
#         .::'         ::::::::::::::'         ``::::.
#     ...:::           ::::::::::::'              ``::.
#    ````':.          ':::::::::'                  ::::..
#                       '.:::::'                    ':'````..
# ------------------------------------------------------------
import pathlib
from docx import Document


class DocxBase(object):
    """
    文档处理基类
    """

    def __init__(self, docx_file):
        self.doc = Document(docx_file)
        self.level = {
            2: "紧急",
            3: "高危",
            4: "中危",
            5: "低危",
            6: "信息",
        }

    def get_dbapp_url(self):
        url = ""
        level = 6
        for table in self.doc.tables:
            title = table.cell(0, 0).text.strip()
            if title == "网站URL" and len(table.rows) > 1 and len(table.columns) == 8:
                url = table.cell(1, 0).text.strip()
                for i in range(2, 7):
                    if table.cell(1, i).text.strip() != "0":
                        level = i
                        break

        return self.level[level], url


if __name__ == "__main__":
    docx_dir = '/Users/hypdncy/Desktop/2.无漏洞/'
    docx_files = pathlib.Path(docx_dir).glob("*.docx")
    for file in docx_files:
        print(file.resolve())
        if file.name.startswith("~"):
            continue
        level, url = DocxBase(file.resolve()).get_dbapp_url()
        if url:
            file_name = url.replace("://", "_").replace(":", "_").replace("/", "_")
            if file_name.endswith("_"):
                file_name = file_name[:-1]
            file_name = level + "_" + file_name + ".docx"
            file.rename(pathlib.PurePath(docx_dir, file_name))
